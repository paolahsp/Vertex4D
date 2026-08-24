from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as PdfImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "demo_rehearsal_assets"
MANIFEST = ASSETS / "capture_manifest.json"
OUT_DOCX = ROOT / "docs" / "VERTEX_DEMO_REHEARSAL_REPORT.docx"
OUT_PDF = ROOT / "docs" / "VERTEX_DEMO_REHEARSAL_REPORT.pdf"

COHORT_ID = "cohort_83a643a003"
BRANCH = "agent/vertex-v4-redesign"
START_COMMIT = "6709b1e"
DEMO_EMAIL = "facilitator@northstar-demo.example"
DEMO_PASSWORD = "vertex-demo-2026"
BASE_URL = "http://127.0.0.1:8010"


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(hdr[idx], header, bold=True)
        set_cell_fill(hdr[idx], "F2F4F7")
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if widths:
            hdr[idx].width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cells[idx].width = Inches(widths[idx])
    return table


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.1)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def add_status_run(paragraph, label: str, color: str) -> None:
    run = paragraph.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(color)


def add_fit_picture(doc: Document, path: Path, caption: str) -> None:
    image = Image.open(path)
    width_px, height_px = image.size
    max_width = 6.35
    max_height = 6.35
    aspect = width_px / height_px
    if width_px >= height_px:
        width = min(max_width, max_height * aspect)
        doc.add_picture(str(path), width=Inches(width))
    else:
        height = min(max_height, max_width / aspect)
        doc.add_picture(str(path), height=Inches(height))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(85, 85, 85)


def load_font(size: int, bold: bool = False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def rounded_rect(draw: ImageDraw.ImageDraw, box, fill, outline, radius=16, width=2) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw: ImageDraw.ImageDraw, box, text, font, fill=(20, 24, 33), spacing=4) -> None:
    lines = text.split("\n")
    sizes = [draw.textbbox((0, 0), line, font=font) for line in lines]
    heights = [bbox[3] - bbox[1] for bbox in sizes]
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = box[1] + ((box[3] - box[1]) - total_h) / 2
    for line, bbox, h in zip(lines, sizes, heights):
        w = bbox[2] - bbox[0]
        x = box[0] + ((box[2] - box[0]) - w) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += h + spacing


def draw_arrow(draw: ImageDraw.ImageDraw, start, end, color=(102, 75, 255), width=4) -> None:
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        points = [(x2, y2), (x2 - 14, y2 - 8), (x2 - 14, y2 + 8)]
    else:
        points = [(x2, y2), (x2 + 14, y2 - 8), (x2 + 14, y2 + 8)]
    draw.polygon(points, fill=color)


def make_sequence_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(img)
    title = load_font(42, True)
    font = load_font(25, True)
    small = load_font(18)
    draw.text((70, 50), "VERTEX 4D demo rehearsal sequence", font=title, fill=(17, 24, 39))
    steps = [
        ("1", "Login", "Facilitator auth"),
        ("2", "Dashboard", "Program metrics"),
        ("3", "Cohort", "Case room"),
        ("4", "Queue", "Intervention needs"),
        ("5", "Case", "Comments + rubric"),
        ("6", "Memo", "Decision trace"),
        ("7", "Memo PDF", "Browser print"),
        ("8", "Outcome", "Cohort report"),
        ("9", "Report PDF", "Buyer artifact"),
    ]
    boxes = []
    for i, step in enumerate(steps):
        if i < 5:
            x = 120
            y = 145 + i * 165
        else:
            x = 1000
            y = 230 + (i - 5) * 165
        box = (x, y, x + 520, y + 112)
        fill = (245, 247, 255) if i < 5 else (244, 255, 249)
        rounded_rect(draw, box, fill, (174, 157, 255))
        draw.text((x + 22, y + 18), step[0], font=font, fill=(102, 75, 255))
        draw.text((x + 70, y + 18), step[1], font=font, fill=(17, 24, 39))
        draw.text((x + 70, y + 62), step[2], font=small, fill=(75, 85, 99))
        boxes.append(box)
    for i in range(len(boxes) - 1):
        b1, b2 = boxes[i], boxes[i + 1]
        if i == 4:
            y_mid = (b1[1] + b1[3]) // 2
            draw_arrow(draw, (b1[2] + 24, y_mid), (b2[0] - 24, (b2[1] + b2[3]) // 2))
        else:
            draw_arrow(draw, ((b1[0] + b1[2]) // 2, b1[3] + 10), ((b2[0] + b2[2]) // 2, b2[1] - 10))
    img.save(path)


def make_dual_flow_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(img)
    title = load_font(42, True)
    lane = load_font(30, True)
    font = load_font(23, True)
    small = load_font(18)
    draw.text((70, 50), "Current dual flow", font=title, fill=(17, 24, 39))
    draw.text((90, 145), "Founder flow", font=lane, fill=(17, 24, 39))
    draw.text((90, 540), "Facilitator / institution flow", font=lane, fill=(17, 24, 39))
    founder = [
        ("Decision Case", "Business question"),
        ("Locked Baseline", "Initial intuition"),
        ("Golden Path", "Evidence artifacts"),
        ("DecisionRecord", "Final decision"),
        ("Decision Memo", "Traceable artifact"),
    ]
    facilitator = [
        ("Cohort", "Institution container"),
        ("Case management", "Runs + teams"),
        ("Intervention queue", "Missing data + risk"),
        ("Comments", "Facilitator guidance"),
        ("Rubric scores", "Pre/post quality"),
        ("Outcome Report", "Cohort evidence"),
    ]
    rows = [(founder, 205, (245, 247, 255)), (facilitator, 600, (244, 255, 249))]
    for items, y, fill in rows:
        boxes = []
        w = 250 if len(items) == 6 else 295
        gap = 25
        x = 90
        for title_text, detail in items:
            box = (x, y, x + w, y + 120)
            rounded_rect(draw, box, fill, (174, 157, 255))
            center_text(draw, (x + 10, y + 14, x + w - 10, y + 62), title_text, font)
            center_text(draw, (x + 10, y + 70, x + w - 10, y + 106), detail, small, fill=(75, 85, 99))
            boxes.append(box)
            x += w + gap
        for i in range(len(boxes) - 1):
            b1, b2 = boxes[i], boxes[i + 1]
            draw_arrow(draw, (b1[2] + 6, (b1[1] + b1[3]) // 2), (b2[0] - 6, (b2[1] + b2[3]) // 2))
    draw.line([(965, 325), (965, 600)], fill=(102, 75, 255), width=4)
    draw.polygon([(965, 600), (953, 580), (977, 580)], fill=(102, 75, 255))
    draw.text((995, 432), "Decision evidence feeds cohort oversight", font=small, fill=(75, 85, 99))
    img.save(path)


def apply_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_metadata(doc: Document) -> None:
    rows = [
        ["Fecha de ensayo", "2026-08-24"],
        ["Branch", BRANCH],
        ["Commit base del ensayo", START_COMMIT],
        ["Demo cohort", COHORT_ID],
        ["Facilitator login", f"{DEMO_EMAIL} / {DEMO_PASSWORD}"],
        ["Servidor local usado", BASE_URL],
    ]
    add_table(doc, ["Campo", "Valor"], rows, widths=[1.8, 4.7])


def add_roadmap(doc: Document) -> None:
    doc.add_heading("Estado frente al roadmap", level=1)
    p = doc.add_paragraph()
    p.add_run("Determinacion: ").bold = True
    p.add_run(
        "VERTEX esta en cierre de Meses 0-2 y entrando a Meses 2-4. "
        "Tiene algunos artifacts de Meses 4-6 adelantados, especialmente Outcome Report y print/export por navegador, "
        "pero aun falta validacion comercial real con cohorts pagados."
    )
    rows = [
        [
            "Meses 0-2: Pilot hardening",
            "Mayormente logrado",
            "Journey demo estable, baseline/post assessment, Decision Memo, dashboard facilitador, analytics/eventos y reportes print.",
            "Ajuste de contraste/responsive QA completo; Evidence Gate mas explicito como gate de producto.",
        ],
        [
            "Meses 2-4: Alpha y Beta pilots",
            "Listo para iniciar",
            "Seed demo cubre 5 casos, intervention queue, comentarios, rubric pre/post y pricing experiments como narrativa.",
            "Ejecutar 8-12 founders reales, luego 20-40 equipos; medir friccion y completar feedback real.",
        ],
        [
            "Meses 4-6: Commercial V1",
            "Parcial adelantado",
            "Outcome reports, print artifacts, checklist, proposal, follow-up docs y roles basicos.",
            "Cohort setup mas robusto, better export, permisos formales, onboarding institucional y 3-5 pilotos pagados.",
        ],
        [
            "Meses 6-9: Institutional readiness",
            "No iniciado",
            "Hay privacidad practica y docs iniciales, pero no enterprise readiness.",
            "Multi-program dashboard, SSO/SAML, security documentation, white-label ligero, export/API basico.",
        ],
        [
            "Meses 9-12: Scale test",
            "No iniciado",
            "La base conceptual esta, pero aun no hay escala comercial.",
            "Annual licenses, self-service paid, vertical templates, cohort benchmarking, LMS integrations si ventas lo justifica.",
        ],
    ]
    add_table(doc, ["Roadmap", "Estado", "Evidencia actual", "Falta"], rows, widths=[1.5, 1.1, 2.0, 1.9])


def build_doc() -> None:
    with MANIFEST.open("r", encoding="utf-8") as f:
        steps = json.load(f)

    sequence = ASSETS / "flow_sequence.png"
    dual = ASSETS / "dual_flow.png"
    make_sequence_diagram(sequence)
    make_dual_flow_diagram(dual)

    doc = Document()
    apply_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("VERTEX 4D - Demo Rehearsal Evidence Pack")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(17, 24, 39)

    subtitle = doc.add_paragraph(
        "Ensayo completo con seed demo: dashboard, cohort, intervention queue, comments/rubric, Decision Memo print y Outcome Report print."
    )
    subtitle.runs[0].font.color.rgb = RGBColor(75, 85, 99)
    add_metadata(doc)

    doc.add_heading("Resultado ejecutivo", level=1)
    for item in [
        "El flujo buyer-facing completo se pudo recorrer de punta a punta con datos demo realistas.",
        "Se generaron PDFs de navegador para el Decision Memo y el Cohort Outcome Report.",
        "Las vistas print son las mas presentables para entregar al buyer: limpias, trazables y con disclaimer de no-overclaim.",
        "La UI web normal funciona, pero el tema claro tiene bajo contraste en tablas, labels secundarios y algunos bloques largos.",
        "El producto esta listo para un rehearsal interno y casi listo para demo buyer; antes de buyer en vivo conviene corregir contraste y probar responsive mobile.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Artifacts generados", level=1)
    add_table(
        doc,
        ["Artifact", "Ruta"],
        [
            ["Decision Memo PDF", "docs/demo_rehearsal_assets/decision-memo-run_demo_economics_changed.pdf"],
            ["Cohort Outcome Report PDF", "docs/demo_rehearsal_assets/cohort-outcome-report-demo.pdf"],
            ["Capturas PNG", "docs/demo_rehearsal_assets/01-login.png a 09-outcome-report-print.png"],
            ["Manifest del ensayo", "docs/demo_rehearsal_assets/capture_manifest.json"],
        ],
        widths=[2.2, 4.3],
    )

    doc.add_heading("Secuencia en grafica", level=1)
    doc.add_paragraph("La ruta ejecutada replica lo que veria un buyer al observar una demo institucional guiada.")
    add_fit_picture(doc, sequence, "Grafica de la secuencia recorrida durante el ensayo.")

    doc.add_heading("Flujo dual actual", level=1)
    doc.add_paragraph(
        "VERTEX ya opera con dos carriles conectados: el founder produce evidencia y decision traceability; el facilitator convierte esa evidencia en seguimiento de cohort, intervenciones y artifacts institucionales."
    )
    add_fit_picture(doc, dual, "Flujo dual actual: founder flow y facilitator/institution flow.")

    doc.add_heading("Capturas y resultado de cada paso", level=1)
    rows = []
    for idx, step in enumerate(steps, 1):
        rows.append([str(idx), step["title"], step["description"], step["result"], step["url"]])
    add_table(doc, ["Paso", "Pantalla", "Que hice", "Que paso", "URL"], rows, widths=[0.45, 1.2, 1.75, 1.8, 1.3])

    for idx, step in enumerate(steps, 1):
        doc.add_page_break()
        doc.add_heading(f"Paso {idx}: {step['title']}", level=2)
        doc.add_paragraph(f"Que hice: {step['description']}")
        doc.add_paragraph(f"Que paso: {step['result']}")
        doc.add_paragraph(f"URL: {step['url']}")
        add_fit_picture(doc, ROOT / step["screenshot"], f"Captura {idx}: {step['title']}")

    doc.add_page_break()
    add_roadmap(doc)

    doc.add_heading("Faltantes recomendados antes de buyer demo", level=1)
    for item in [
        "Ajustar contraste del tema claro en dashboard, cohort table, report web y texto secundario.",
        "Hacer QA responsive mobile real, porque aun no se valido en este ensayo.",
        "Pulir el layout web normal del Decision Memo para que el bloque de decision final no domine excesivamente.",
        "Convertir Evidence Gate en un checkpoint explicito dentro del journey, no solo evidencia trazable en artifacts.",
        "Ejecutar un rehearsal con narracion de 8-12 minutos usando docs/VERTEX_DEMO_SCRIPT.md y medir tiempos.",
        "Despues del demo interno, correr 8-12 founders reales como Alpha pilot y reemplazar demo assumptions por evidencia real.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Handoff operativo", level=1)
    add_table(
        doc,
        ["Tema", "Estado actual"],
        [
            ["Seed command", "python -m scripts.seed_demo_cohort --json"],
            ["Reset command", "python -m scripts.seed_demo_cohort --reset-demo"],
            ["Start app", "python -m uvicorn main:app --host 127.0.0.1 --port 8010"],
            ["Facilitator", f"{DEMO_EMAIL} / {DEMO_PASSWORD}"],
            ["Dashboard", f"{BASE_URL}/dashboard/facilitator"],
            ["Cohort", f"{BASE_URL}/dashboard/facilitator/cohorts/{COHORT_ID}"],
            ["Decision Memo print", f"{BASE_URL}/dashboard/lab/decision-memo/print?run_id=run_demo_economics_changed"],
            ["Outcome Report print", f"{BASE_URL}/dashboard/facilitator/cohorts/{COHORT_ID}/outcome-report/print"],
            ["No push", "Confirmado: no se hizo push durante este ensayo."],
        ],
        widths=[1.7, 4.8],
    )

    doc.add_heading("Notas de verdad comercial", level=1)
    for item in [
        "El demo prueba decision quality, traceability, before/after evidence e intervention need.",
        "No prueba market success, investment readiness, startup validation garantizada ni due diligence del buyer.",
        "Los datos son demo y estan marcados como Northstar / Demo; no deben mezclarse con claims de pilotos reales.",
        "El siguiente paso recomendado es un rehearsal buyer-facing cronometrado y despues un Alpha pilot con 8-12 founders reales.",
    ]:
        add_bullet(doc, item)

    doc.core_properties.title = "VERTEX 4D - Demo Rehearsal Evidence Pack"
    doc.core_properties.author = "Codex"
    doc.core_properties.created = datetime.now(timezone.utc)
    doc.save(OUT_DOCX)


def pdf_table(headers: list[str], rows: list[list[str]], col_widths: list[float]):
    data = [[Paragraph(f"<b>{h}</b>", PDF_STYLES["TableCell"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(cell).replace("\n", "<br/>"), PDF_STYLES["TableCell"]) for cell in row])
    table = Table(data, colWidths=[w * inch for w in col_widths], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D0D5DD")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def pdf_fit_image(path: Path, caption: str, max_width: float = 6.5 * inch, max_height: float = 5.6 * inch):
    image = Image.open(path)
    width_px, height_px = image.size
    ratio = min(max_width / width_px, max_height / height_px)
    width = width_px * ratio
    height = height_px * ratio
    return [
        PdfImage(str(path), width=width, height=height),
        Paragraph(f"<i>{caption}</i>", PDF_STYLES["Caption"]),
    ]


def build_pdf() -> None:
    with MANIFEST.open("r", encoding="utf-8") as f:
        steps = json.load(f)

    sequence = ASSETS / "flow_sequence.png"
    dual = ASSETS / "dual_flow.png"
    make_sequence_diagram(sequence)
    make_dual_flow_diagram(dual)

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title="VERTEX 4D - Demo Rehearsal Evidence Pack",
        author="Codex",
    )
    story = []
    story.append(Paragraph("VERTEX 4D - Demo Rehearsal Evidence Pack", PDF_STYLES["Title"]))
    story.append(
        Paragraph(
            "Ensayo completo con seed demo: dashboard, cohort, intervention queue, comments/rubric, Decision Memo print y Outcome Report print.",
            PDF_STYLES["Body"],
        )
    )
    story.append(Spacer(1, 0.12 * inch))
    story.append(
        pdf_table(
            ["Campo", "Valor"],
            [
                ["Fecha de ensayo", "2026-08-24"],
                ["Branch", BRANCH],
                ["Commit base del ensayo", START_COMMIT],
                ["Demo cohort", COHORT_ID],
                ["Facilitator login", f"{DEMO_EMAIL} / {DEMO_PASSWORD}"],
                ["Servidor local usado", BASE_URL],
            ],
            [1.8, 4.7],
        )
    )
    story.append(Spacer(1, 0.16 * inch))
    story.append(Paragraph("Resultado ejecutivo", PDF_STYLES["H1"]))
    for item in [
        "El flujo buyer-facing completo se pudo recorrer de punta a punta con datos demo realistas.",
        "Se generaron PDFs de navegador para el Decision Memo y el Cohort Outcome Report.",
        "Las vistas print son las mas presentables para entregar al buyer: limpias, trazables y con disclaimer de no-overclaim.",
        "La UI web normal funciona, pero el tema claro tiene bajo contraste en tablas, labels secundarios y algunos bloques largos.",
        "El producto esta listo para un rehearsal interno y casi listo para demo buyer; antes de buyer en vivo conviene corregir contraste y probar responsive mobile.",
    ]:
        story.append(Paragraph(f"- {item}", PDF_STYLES["Body"]))

    story.append(Paragraph("Artifacts generados", PDF_STYLES["H1"]))
    story.append(
        pdf_table(
            ["Artifact", "Ruta"],
            [
                ["Decision Memo PDF", "docs/demo_rehearsal_assets/decision-memo-run_demo_economics_changed.pdf"],
                ["Cohort Outcome Report PDF", "docs/demo_rehearsal_assets/cohort-outcome-report-demo.pdf"],
                ["Capturas PNG", "docs/demo_rehearsal_assets/01-login.png a 09-outcome-report-print.png"],
                ["Manifest del ensayo", "docs/demo_rehearsal_assets/capture_manifest.json"],
            ],
            [2.0, 4.5],
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("Secuencia en grafica", PDF_STYLES["H1"]))
    story.append(Paragraph("La ruta ejecutada replica lo que veria un buyer al observar una demo institucional guiada.", PDF_STYLES["Body"]))
    story.extend(pdf_fit_image(sequence, "Grafica de la secuencia recorrida durante el ensayo.", max_height=4.9 * inch))
    story.append(Paragraph("Flujo dual actual", PDF_STYLES["H1"]))
    story.append(
        Paragraph(
            "VERTEX ya opera con dos carriles conectados: el founder produce evidencia y decision traceability; el facilitator convierte esa evidencia en seguimiento de cohort, intervenciones y artifacts institucionales.",
            PDF_STYLES["Body"],
        )
    )
    story.extend(pdf_fit_image(dual, "Flujo dual actual: founder flow y facilitator/institution flow.", max_height=4.2 * inch))

    story.append(PageBreak())
    story.append(Paragraph("Capturas y resultado de cada paso", PDF_STYLES["H1"]))
    story.append(
        pdf_table(
            ["Paso", "Pantalla", "Que hice", "Que paso"],
            [[str(i), step["title"], step["description"], step["result"]] for i, step in enumerate(steps, 1)],
            [0.42, 1.25, 2.1, 2.73],
        )
    )

    for idx, step in enumerate(steps, 1):
        story.append(PageBreak())
        story.append(Paragraph(f"Paso {idx}: {step['title']}", PDF_STYLES["H1"]))
        story.append(Paragraph(f"<b>Que hice:</b> {step['description']}", PDF_STYLES["Body"]))
        story.append(Paragraph(f"<b>Que paso:</b> {step['result']}", PDF_STYLES["Body"]))
        story.append(Paragraph(f"<b>URL:</b> {step['url']}", PDF_STYLES["Small"]))
        story.extend(pdf_fit_image(ROOT / step["screenshot"], f"Captura {idx}: {step['title']}"))

    story.append(PageBreak())
    story.append(Paragraph("Estado frente al roadmap", PDF_STYLES["H1"]))
    story.append(
        Paragraph(
            "<b>Determinacion:</b> VERTEX esta en cierre de Meses 0-2 y entrando a Meses 2-4. "
            "Tiene algunos artifacts de Meses 4-6 adelantados, especialmente Outcome Report y print/export por navegador, "
            "pero aun falta validacion comercial real con cohorts pagados.",
            PDF_STYLES["Body"],
        )
    )
    story.append(
        pdf_table(
            ["Roadmap", "Estado", "Evidencia actual", "Falta"],
            [
                [
                    "Meses 0-2: Pilot hardening",
                    "Mayormente logrado",
                    "Journey demo estable, baseline/post assessment, Decision Memo, dashboard facilitador, analytics/eventos y reportes print.",
                    "Ajuste de contraste/responsive QA completo; Evidence Gate mas explicito como gate de producto.",
                ],
                [
                    "Meses 2-4: Alpha y Beta pilots",
                    "Listo para iniciar",
                    "Seed demo cubre 5 casos, intervention queue, comentarios, rubric pre/post y pricing experiments como narrativa.",
                    "Ejecutar 8-12 founders reales, luego 20-40 equipos; medir friccion y completar feedback real.",
                ],
                [
                    "Meses 4-6: Commercial V1",
                    "Parcial adelantado",
                    "Outcome reports, print artifacts, checklist, proposal, follow-up docs y roles basicos.",
                    "Cohort setup mas robusto, better export, permisos formales, onboarding institucional y 3-5 pilotos pagados.",
                ],
                [
                    "Meses 6-9: Institutional readiness",
                    "No iniciado",
                    "Hay privacidad practica y docs iniciales, pero no enterprise readiness.",
                    "Multi-program dashboard, SSO/SAML, security documentation, white-label ligero, export/API basico.",
                ],
                [
                    "Meses 9-12: Scale test",
                    "No iniciado",
                    "La base conceptual esta, pero aun no hay escala comercial.",
                    "Annual licenses, self-service paid, vertical templates, cohort benchmarking, LMS integrations si ventas lo justifica.",
                ],
            ],
            [1.35, 0.95, 2.05, 2.15],
        )
    )

    story.append(Paragraph("Faltantes recomendados antes de buyer demo", PDF_STYLES["H1"]))
    for item in [
        "Ajustar contraste del tema claro en dashboard, cohort table, report web y texto secundario.",
        "Hacer QA responsive mobile real, porque aun no se valido en este ensayo.",
        "Pulir el layout web normal del Decision Memo para que el bloque de decision final no domine excesivamente.",
        "Convertir Evidence Gate en un checkpoint explicito dentro del journey, no solo evidencia trazable en artifacts.",
        "Ejecutar un rehearsal con narracion de 8-12 minutos usando docs/VERTEX_DEMO_SCRIPT.md y medir tiempos.",
        "Despues del demo interno, correr 8-12 founders reales como Alpha pilot y reemplazar demo assumptions por evidencia real.",
    ]:
        story.append(Paragraph(f"- {item}", PDF_STYLES["Body"]))

    story.append(Paragraph("Handoff operativo", PDF_STYLES["H1"]))
    story.append(
        pdf_table(
            ["Tema", "Estado actual"],
            [
                ["Seed command", "python -m scripts.seed_demo_cohort --json"],
                ["Reset command", "python -m scripts.seed_demo_cohort --reset-demo"],
                ["Start app", "python -m uvicorn main:app --host 127.0.0.1 --port 8010"],
                ["Facilitator", f"{DEMO_EMAIL} / {DEMO_PASSWORD}"],
                ["Dashboard", f"{BASE_URL}/dashboard/facilitator"],
                ["Cohort", f"{BASE_URL}/dashboard/facilitator/cohorts/{COHORT_ID}"],
                ["Decision Memo print", f"{BASE_URL}/dashboard/lab/decision-memo/print?run_id=run_demo_economics_changed"],
                ["Outcome Report print", f"{BASE_URL}/dashboard/facilitator/cohorts/{COHORT_ID}/outcome-report/print"],
                ["No push", "Confirmado: no se hizo push durante este ensayo."],
            ],
            [1.65, 4.85],
        )
    )
    story.append(Paragraph("Notas de verdad comercial", PDF_STYLES["H1"]))
    for item in [
        "El demo prueba decision quality, traceability, before/after evidence e intervention need.",
        "No prueba market success, investment readiness, startup validation garantizada ni due diligence del buyer.",
        "Los datos son demo y estan marcados como Northstar / Demo; no deben mezclarse con claims de pilotos reales.",
        "El siguiente paso recomendado es un rehearsal buyer-facing cronometrado y despues un Alpha pilot con 8-12 founders reales.",
    ]:
        story.append(Paragraph(f"- {item}", PDF_STYLES["Body"]))

    doc.build(story)


def make_pdf_styles():
    base = getSampleStyleSheet()
    return {
        "Title": ParagraphStyle(
            "VertexTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=21,
            leading=25,
            alignment=0,
            textColor=colors.HexColor("#111827"),
            spaceAfter=10,
        ),
        "H1": ParagraphStyle(
            "VertexH1",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=12,
            spaceAfter=7,
        ),
        "Body": ParagraphStyle(
            "VertexBody",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=13,
            spaceAfter=6,
        ),
        "Small": ParagraphStyle(
            "VertexSmall",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#555555"),
            spaceAfter=6,
        ),
        "Caption": ParagraphStyle(
            "VertexCaption",
            parent=base["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#555555"),
            alignment=TA_CENTER,
            spaceBefore=4,
            spaceAfter=6,
        ),
        "TableCell": ParagraphStyle(
            "VertexTableCell",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.2,
            leading=9,
            spaceAfter=0,
        ),
    }


PDF_STYLES = make_pdf_styles()


if __name__ == "__main__":
    build_doc()
    build_pdf()
